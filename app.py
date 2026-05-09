import io
from datetime import datetime
from flask import Flask, jsonify, render_template, request, send_file
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fpdf import FPDF

app = Flask(__name__)


def format_date_for_export(date_str):
    if not date_str:
        return ''
    try:
        return datetime.strptime(date_str, '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        return date_str


def build_export_filename(date_str, time_str, extension):
    now = datetime.now()
    date_part = now.strftime('%d%m%Y')
    time_part = now.strftime('%H%M')

    if date_str:
        try:
            date_part = datetime.strptime(date_str, '%Y-%m-%d').strftime('%d%m%Y')
        except ValueError:
            pass

    if time_str:
        try:
            time_part = datetime.strptime(time_str, '%H:%M').strftime('%H%M')
        except ValueError:
            pass

    return f'mom_{date_part}_{time_part}.{extension}'


def build_word_bytes(date, time, subject, attendees, points, action_items):
    doc = Document()
    valid_action_items = [item for item in action_items if item.get('desc', '').strip()]

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Verdana'
    normal_style.font.size = Pt(10)
    normal_style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    run = title.add_run('Minutes of Meeting')
    run.font.name = 'Verdana'
    run.font.size = Pt(11)
    run.font.bold = True

    if subject:
        subject_para = doc.add_paragraph()
        subject_run = subject_para.add_run(f'Subject: {subject}')
        subject_run.font.name = 'Verdana'
        subject_run.font.size = Pt(11)
        subject_run.font.bold = True

    doc.add_paragraph(f'Date: {format_date_for_export(date)}')
    doc.add_paragraph(f'Time: {time}')
    doc.add_paragraph()

    heading = doc.add_paragraph()
    h_run = heading.add_run('Members Present')
    h_run.font.name = 'Verdana'
    h_run.font.size = Pt(11)
    h_run.bold = True
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(8)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'S.No'
    hdr_cells[1].text = 'Name'
    hdr_cells[2].text = 'Designation / Role'
    for cell in hdr_cells:
        shading = parse_xml(r'<w:shd {} w:fill="D9EAF7"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)

    sno = 1
    for attendee in attendees:
        name = attendee.get('name', '').strip()
        role = attendee.get('role', '').strip()
        if not name:
            continue
        row_cells = table.add_row().cells
        row_cells[0].text = str(sno)
        row_cells[1].text = name
        row_cells[2].text = role
        sno += 1
    doc.add_paragraph()

    heading = doc.add_paragraph()
    h_run = heading.add_run('Points Discussed')
    h_run.font.name = 'Verdana'
    h_run.font.size = Pt(11)
    h_run.bold = True
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(8)
    points_para = doc.add_paragraph(points if points else '')
    points_para.paragraph_format.space_after = Pt(8)

    if valid_action_items:
        heading = doc.add_paragraph()
        h_run = heading.add_run('Action Items')
        h_run.font.name = 'Verdana'
        h_run.font.size = Pt(11)
        h_run.bold = True
        heading.paragraph_format.space_before = Pt(14)
        heading.paragraph_format.space_after = Pt(8)

        action_table = doc.add_table(rows=1, cols=3)
        action_table.style = 'Table Grid'
        action_hdr = action_table.rows[0].cells
        action_hdr[0].text = 'Action Item'
        action_hdr[1].text = 'Responsibility'
        action_hdr[2].text = 'Timeline'
        for cell in action_hdr:
            shading = parse_xml(r'<w:shd {} w:fill="D9EAF7"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)

        for item in valid_action_items:
            desc = item.get('desc', '').strip()
            action_row = action_table.add_row().cells
            action_row[0].text = desc
            action_row[1].text = item.get('resp', '').strip()
            action_row[2].text = item.get('date', '').strip()
        doc.add_paragraph()

    heading = doc.add_paragraph()
    h_run = heading.add_run('Signatures')
    h_run.font.name = 'Verdana'
    h_run.font.size = Pt(11)
    h_run.bold = True
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(8)
    para = doc.add_paragraph('Agreed and acknowledged by the members present:')
    para.runs[0].italic = True

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.autofit = False
    sig_table.allow_autofit = False

    written = 0
    for attendee in attendees:
        name = attendee.get('name', '').strip()
        role = attendee.get('role', '').strip()
        if not name:
            continue
        if written % 2 == 0:
            row = sig_table.rows[0] if written == 0 else sig_table.add_row()
        else:
            row = sig_table.rows[-1]
        cell = row.cells[written % 2]
        p = cell.paragraphs[0]
        p.alignment = 1
        r = p.add_run('_' * 20 + '\n' + name + ('\n' + role if role else ''))
        r.font.name = 'Verdana'
        r.font.size = Pt(10)
        written += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_pdf_bytes(date, time, subject, attendees, points, action_items):
    pdf = FPDF()
    valid_action_items = [item for item in action_items if item.get('desc', '').strip()]
    valid_attendees = [a for a in attendees if a.get('name', '').strip()]
    font_family = 'Arial'
    try:
        pdf.add_font('Verdana', '', 'C:/Windows/Fonts/verdana.ttf', uni=True)
        pdf.add_font('Verdana', 'B', 'C:/Windows/Fonts/verdanab.ttf', uni=True)
        pdf.add_font('Verdana', 'I', 'C:/Windows/Fonts/verdanai.ttf', uni=True)
        font_family = 'Verdana'
    except Exception:
        font_family = 'Arial'

    pdf.add_page()
    pdf.set_font(font_family, 'B', 11)
    pdf.cell(0, 10, txt='Minutes of Meeting', ln=True, align='C')
    if subject:
        pdf.multi_cell(0, 6, txt=f'Subject: {subject}', align='C')
    pdf.set_font(font_family, '', 10)
    pdf.ln(3)
    pdf.cell(0, 8, txt=f'Date: {format_date_for_export(date)}', ln=True)
    pdf.cell(0, 8, txt=f'Time: {time}', ln=True)
    pdf.ln(4)

    pdf.set_font(font_family, 'B', 11)
    pdf.cell(0, 8, txt='Members Present:', ln=True)
    pdf.set_font(font_family, '', 10)
    widths = [15, 80, 95]
    pdf.set_fill_color(217, 234, 247)
    pdf.cell(widths[0], 7, 'S.No', border=1, align='C', fill=True)
    pdf.cell(widths[1], 7, 'Name', border=1, align='C', fill=True)
    pdf.cell(widths[2], 7, 'Designation / Role', border=1, align='C', ln=True, fill=True)

    sno = 1
    for attendee in attendees:
        name = attendee.get('name', '').strip()
        role = attendee.get('role', '').strip()
        if not name:
            continue
        pdf.cell(widths[0], 7, str(sno), border=1, align='C')
        pdf.cell(widths[1], 7, name[:30], border=1)
        pdf.cell(widths[2], 7, role[:35], border=1, ln=True)
        sno += 1
    pdf.ln(4)

    pdf.set_font(font_family, 'B', 11)
    pdf.cell(0, 8, txt='Points Discussed:', ln=True)
    pdf.set_font(font_family, '', 10)
    pdf.multi_cell(0, 5, points if points else '')
    pdf.ln(4)

    if valid_action_items:
        pdf.set_font(font_family, 'B', 11)
        pdf.cell(0, 8, txt='Action Items:', ln=True)
        pdf.set_font(font_family, '', 10)
        aw = [65, 65, 40]
        pdf.set_fill_color(217, 234, 247)
        pdf.cell(aw[0], 7, 'Action Item', border=1, align='C', fill=True)
        pdf.cell(aw[1], 7, 'Responsibility', border=1, align='C', fill=True)
        pdf.cell(aw[2], 7, 'Timeline', border=1, align='C', ln=True, fill=True)
        for item in valid_action_items:
            desc = item.get('desc', '').strip()
            pdf.cell(aw[0], 7, desc[:35], border=1)
            pdf.cell(aw[1], 7, item.get('resp', '').strip()[:30], border=1)
            pdf.cell(aw[2], 7, item.get('date', '').strip()[:18], border=1, ln=True)
        pdf.ln(4)

    # Signatures section
    pdf.set_font(font_family, 'B', 11)
    pdf.cell(0, 8, txt='Signatures:', ln=True)
    pdf.set_font(font_family, 'I', 10)
    pdf.cell(0, 6, txt='Agreed and acknowledged by the members present:', ln=True)
    pdf.ln(3)
    pdf.set_font(font_family, '', 10)

    sig_per_row = 2
    sig_width = 90
    sig_x_start = 15
    sig_y = pdf.get_y()
    written = 0

    for attendee in valid_attendees:
        col_idx = written % sig_per_row
        row_idx = written // sig_per_row
        x_pos = sig_x_start + (col_idx * sig_width)
        y_pos = sig_y + (row_idx * 24)

        pdf.set_xy(x_pos, y_pos)
        pdf.cell(sig_width, 2, '_' * 30, border=0, align='C')
        pdf.set_xy(x_pos, y_pos + 3)
        pdf.cell(sig_width, 5, attendee.get('name', '').strip(), border=0, align='C')
        pdf.set_xy(x_pos, y_pos + 8)
        pdf.cell(sig_width, 4, attendee.get('role', '').strip(), border=0, align='C')
        written += 1

    out = pdf.output(dest='S').encode('latin-1')
    return io.BytesIO(out)


@app.route('/', methods=['GET'])
def index():
    return render_template('index.html')


@app.route('/api/export/word', methods=['POST'])
def export_word():
    payload = request.get_json(silent=True) or {}
    date = payload.get('date', '')
    time = payload.get('time', '')
    file_buffer = build_word_bytes(
        date,
        time,
        payload.get('subject', ''),
        payload.get('attendees', []),
        payload.get('points', ''),
        payload.get('action_items', []),
    )
    filename = build_export_filename(date, time, 'docx')
    return send_file(
        file_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )


@app.route('/api/export/pdf', methods=['POST'])
def export_pdf():
    payload = request.get_json(silent=True) or {}
    date = payload.get('date', '')
    time = payload.get('time', '')
    file_buffer = build_pdf_bytes(
        date,
        time,
        payload.get('subject', ''),
        payload.get('attendees', []),
        payload.get('points', ''),
        payload.get('action_items', []),
    )
    filename = build_export_filename(date, time, 'pdf')
    return send_file(
        file_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf',
    )


@app.route('/api/health', methods=['GET'])
def health():
    return jsonify({'ok': True})


if __name__ == '__main__':
    app.run(debug=True)